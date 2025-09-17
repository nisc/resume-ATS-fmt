#!/usr/bin/env python3
"""
DOCX for ATS formatter
"""

import argparse
import os
import re
import sys
from pathlib import Path
from typing import Dict, Match, Optional, Set, Tuple


def find_and_activate_venv() -> None:
    """Find and activate the virtual environment if needed."""
    # Check if we're already in a virtual environment
    if hasattr(sys, "real_prefix") or (
        hasattr(sys, "base_prefix") and sys.base_prefix != sys.prefix
    ):
        return  # Already in a virtual environment

    # Get the directory where this script is located (resolve symlinks)
    script_path = Path(__file__).resolve()
    script_dir = script_path.parent

    # The project directory is where the script is located
    project_dir = script_dir

    # Check if .venv exists in the project directory
    venv_python = project_dir / ".venv" / "bin" / "python"
    if not venv_python.exists():
        print(f"Error: Virtual environment not found at {venv_python}")
        print(
            "Please ensure you have a .venv directory in the same dir as this script."
        )
        sys.exit(1)

    # Re-execute the script with the virtual environment Python
    if sys.executable != str(venv_python):
        try:
            os.execv(str(venv_python), [str(venv_python)] + sys.argv)
        except OSError as e:
            print(f"Error: Failed to activate virtual environment: {e}")
            sys.exit(1)


# Activate virtual environment if needed
find_and_activate_venv()

try:
    from docx import Document
    from docx.text.paragraph import Paragraph
except ImportError:
    print(
        "Error: python-docx library not found. "
        "Please install it with: pip install python-docx"
    )
    sys.exit(1)


class TitleCaseConverter:
    """Converts ALL CAPS words to Title Case with exceptions."""

    def __init__(self, exceptions_file: Optional[str] = None):
        """Initialize the converter with optional exceptions file."""
        self.exceptions, self.special_formats = self._load_exceptions(exceptions_file)

    def _load_exceptions(
        self, exceptions_file: Optional[str]
    ) -> Tuple[Set[str], Dict[str, str]]:
        """Load exceptions and special formats from file."""
        # Use default .exceptions file if no file specified
        if exceptions_file is None:
            # Look for .exceptions file relative to the script's directory
            script_dir = Path(__file__).resolve().parent
            exceptions_file = str(script_dir / ".exceptions")

        try:
            with open(exceptions_file, "r", encoding="utf-8") as f:
                exceptions = set()
                special_formats = {}

                for line in f:
                    line = line.strip()
                    if not line:
                        continue

                    if "=" in line:
                        # Format: EXCEPTION=FORMATTED
                        exception, formatted = line.split("=", 1)
                        exception = exception.strip().upper()
                        formatted = formatted.strip()
                        exceptions.add(exception)
                        special_formats[exception] = formatted
                    else:
                        # Format: EXCEPTION (use title case)
                        exceptions.add(line.upper())

                return exceptions, special_formats
        except FileNotFoundError:
            if exceptions_file == ".exceptions":
                print(
                    "Warning: Default .exceptions file not found. "
                    "No exceptions will be applied."
                )
                return set(), {}
            else:
                print(
                    f"Warning: Exceptions file '{exceptions_file}' not found. "
                    "No exceptions will be applied."
                )
                return set(), {}

    def _is_all_caps(self, text: str) -> bool:
        """Check if text is all uppercase letters (including Unicode)."""
        return text.isupper() and text.replace(" ", "").isalpha()

    def _convert_to_title_case(self, text: str) -> str:
        """Convert text to title case, handling exceptions."""
        # Check if this exact text is in our exceptions (case-insensitive)
        if text.upper() in self.exceptions:
            return self._apply_exception_format(text.upper())

        # Default title case conversion
        return text.title()

    def _apply_exception_format(self, exception: str) -> str:
        """Apply proper formatting for exception words."""
        return self.special_formats.get(exception, exception.title())

    def convert_text(self, text: str) -> str:
        """Convert ALL CAPS words in text to Title Case and remove brackets."""
        # First, handle ALL CAPS words followed by brackets
        # Pattern: ALL CAPS word followed by optional whitespace and text in brackets
        # (handles nested brackets)
        bracket_pattern = r"\b([A-Z]{2,})\s*\([^()]*(?:\([^()]*\)[^()]*)*\)"

        def replace_caps_with_brackets(match: Match[str]) -> str:
            caps_word = match.group(1)
            if self._is_all_caps(caps_word) and len(caps_word) >= 5:
                return self._convert_to_title_case(caps_word)
            return match.group()  # Return original if not all caps or too short

        # Remove brackets after ALL CAPS words
        text = re.sub(bracket_pattern, replace_caps_with_brackets, text)

        # First, handle known exception phrases that might contain punctuation
        for exception in self.exceptions:
            if len(exception.split()) > 1:  # Multi-word exceptions
                # Create a pattern that matches the exception with flexible spacing
                # Replace spaces with flexible whitespace and punctuation patterns
                escaped_exception = re.escape(exception)
                # Allow flexible spacing and common punctuation between words
                flexible_pattern = escaped_exception.replace(r"\ ", r"\s*[&\-,–—]?\s*")
                exception_pattern = r"\b" + flexible_pattern + r"\b"
                text = re.sub(
                    exception_pattern,
                    lambda m: self._apply_exception_format(exception),
                    text,
                    flags=re.IGNORECASE,
                )

        # Find all ALL CAPS words and their positions
        # Pattern to match ALL CAPS words (at least 2 characters, all letters)
        # We'll use a comprehensive approach by finding all words and checking
        # if they're all caps
        pattern = r"\b\w{2,}\b"
        matches = list(re.finditer(pattern, text))

        # Group consecutive ALL CAPS words
        sequences = []
        current_sequence: list[Match[str]] = []

        for i, match in enumerate(matches):
            word = match.group()
            # Process words that are all caps
            if not self._is_all_caps(word):
                continue

            if not current_sequence:
                current_sequence = [match]
            else:
                # Check if this word is consecutive with the previous one
                # (separated only by whitespace and common punctuation)
                prev_match = current_sequence[-1]
                gap = text[prev_match.end() : match.start()]
                # Allow whitespace and common punctuation like "and", "&", ",", etc.
                # Also allow short connecting words
                gap_clean = gap.strip()
                if (
                    gap.isspace()
                    or gap == ""
                    or gap_clean in ["and", "&", ",", "-", "–", "—"]
                    or gap_clean.lower()
                    in ["and", "or", "of", "the", "in", "on", "at", "to", "for"]
                    or (len(gap_clean) <= 3 and gap_clean.isalpha())
                ):
                    current_sequence.append(match)
                else:
                    # End current sequence and start new one
                    if current_sequence:
                        sequences.append(current_sequence)
                    current_sequence = [match]

        # Add the last sequence if it exists
        if current_sequence:
            sequences.append(current_sequence)

        # Process each sequence
        result = text
        offset = 0  # Track offset due to previous replacements

        for sequence in sequences:
            if len(sequence) == 1:
                # Single word - only convert if >= 5 characters
                match = sequence[0]
                word = match.group()
                if len(word) >= 5:
                    replacement = self._convert_to_title_case(word)
                    start = match.start() + offset
                    end = match.end() + offset
                    result = result[:start] + replacement + result[end:]
                    offset += len(replacement) - len(word)
            else:
                # Multiple words - convert all words in the sequence
                for match in sequence:
                    word = match.group()
                    replacement = self._convert_to_title_case(word)
                    start = match.start() + offset
                    end = match.end() + offset
                    result = result[:start] + replacement + result[end:]
                    offset += len(replacement) - len(word)

        return result

    def _process_paragraph_runs(self, paragraph: Paragraph) -> None:
        """Process paragraph runs and strip all formatting, set font size to 9.5."""
        # Get all text from the paragraph
        full_text = paragraph.text

        if not full_text.strip():
            return

        # Convert the text
        converted_text = self.convert_text(full_text)

        # Always clear and rebuild the paragraph to ensure no formatting
        paragraph.clear()

        # Add the converted text as plain text with font size 9.5
        if converted_text.strip():
            run = paragraph.add_run(converted_text)
            # Set font size to 9.5 points (9.5 * 12700 = 120650 for docx units)
            run.font.size = 120650

    def process_docx(self, input_path: str, output_path: Optional[str] = None) -> str:
        """Process a .docx file and convert ALL CAPS words to Title Case."""
        input_file = Path(input_path)
        if not input_file.exists():
            raise FileNotFoundError(f"Input file not found: {input_path}")

        if input_file.suffix.lower() != ".docx":
            raise ValueError("Input file must be a .docx file")

        # Set output path if not provided
        if output_path is None:
            # Save in the same directory as the input file
            output_path = str(
                input_file.parent / f"ATS_{input_file.stem}{input_file.suffix}"
            )

        # Load the document
        doc = Document(input_path)

        # Process all paragraphs
        for paragraph in doc.paragraphs:
            self._process_paragraph_runs(paragraph)

        # Process all tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._process_paragraph_runs(paragraph)

        # Process headers and footers
        for section in doc.sections:
            # Header
            if section.header:
                for paragraph in section.header.paragraphs:
                    self._process_paragraph_runs(paragraph)

            # Footer
            if section.footer:
                for paragraph in section.footer.paragraphs:
                    self._process_paragraph_runs(paragraph)

        # Save the processed document
        doc.save(output_path)
        return str(output_path)


def main() -> None:
    """Handle command line arguments and process files."""
    parser = argparse.ArgumentParser(
        description="Convert ALL CAPS words in .docx files to Title Case "
        "with exceptions"
    )
    parser.add_argument("input_file", help="Path to the input .docx file")
    parser.add_argument(
        "-o",
        "--output",
        help="Path to the output .docx file (default: adds 'ATS_' prefix)",
    )
    parser.add_argument(
        "-e",
        "--exceptions",
        help="Path to a text file containing one exception per line",
    )
    parser.add_argument(
        "--add-exception",
        action="append",
        help="Add a specific exception word (can be used multiple times)",
    )

    args = parser.parse_args()

    try:
        # Create converter
        converter = TitleCaseConverter(args.exceptions)

        # Add command line exceptions
        if args.add_exception:
            for exception in args.add_exception:
                converter.exceptions.add(exception.upper())

        # Process the file
        output_path = converter.process_docx(args.input_file, args.output)

        print(f"Successfully processed: {args.input_file}")
        print(f"Output saved to: {output_path}")

    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
